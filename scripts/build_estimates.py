"""Build backend/src/estimates.json from Open CEDA 2025 and data/sectors.json.

Open CEDA gives kg CO2e per US dollar of producer price, in 2023 dollars, for 400 sectors in each country.
A shopper pays a purchaser price in today's rupees, so each India factor is converted with the workbook's
own tables:

    kg per rupee = kg per USD (producer, 2023)
                   x purchaser-to-producer ratio      (the shop price includes trade and transport margins)
                   x 100 / sector price index 2025    (2023 = 100)
                   / rupees per USD in 2025

The workbook is downloaded once from the public AWS Open Data bucket into research/ (not committed).
Run: python scripts/build_estimates.py
"""

import json
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "research" / "open-ceda-2025.xlsx"
RULES = ROOT / "data" / "sectors.json"
OUTPUT = ROOT / "backend" / "src" / "estimates.json"

BUCKET = "open-ceda"
BUCKET_REGION = "us-west-2"
KEY = "data/Open CEDA 2025 by Watershed.xlsx"
COUNTRY = "IND"
PRICE_YEAR = "2025"

DATASET = {
    "name": "Open CEDA 2025",
    "attribution": "CEDA by Watershed",
    "licence": "CC BY-SA 4.0",
    "source": "https://registry.opendata.aws/open-ceda/",
    "homepage": "https://openceda.org/",
    "country": "India",
    "baseYear": 2023,
    "priceYear": int(PRICE_YEAR),
}


def text(value):
    return None if value is None else str(value).strip()


def download():
    import boto3
    from botocore import UNSIGNED
    from botocore.config import Config

    WORKBOOK.parent.mkdir(exist_ok=True)
    client = boto3.client("s3", region_name=BUCKET_REGION, config=Config(signature_version=UNSIGNED))
    client.download_file(BUCKET, KEY, str(WORKBOOK))


def sheet_rows(workbook, name):
    return list(workbook[name].iter_rows(values_only=True))


def row_by_code(rows, header_index, values_index):
    """Map sector codes to values. Codes are stored as numbers in some sheets, so compare them as text."""
    return dict(zip(map(text, rows[header_index][1:]), rows[values_index][1:]))


def india_factors(workbook):
    raw = sheet_rows(workbook, "GHG_t_Raw")
    if text(raw[0][1]) != str(DATASET["baseYear"]) or text(raw[1][1]) != "Producer price" or text(raw[1][3]) != "US Dollar":
        sys.exit("GHG_t_Raw no longer holds 2023 producer-price USD factors. Check the workbook before going on.")
    codes = [text(code) for code in raw[3]]
    country = next(row for row in raw if row[0] == COUNTRY)

    conversion_rows = sheet_rows(workbook, "Purchaser - producer conversion")
    conversion = row_by_code(conversion_rows, 4, 5)
    index_rows = sheet_rows(workbook, "Sector level Price Index")
    price_index = row_by_code(index_rows, 4, next(i for i, row in enumerate(index_rows) if text(row[0]) == PRICE_YEAR))

    rates = sheet_rows(workbook, "Exchange rates")
    year_column = list(map(text, rates[3])).index(PRICE_YEAR)
    inr_per_usd = float(next(row for row in rates if row[0] == COUNTRY)[year_column])

    names = {text(row[0]): text(row[1]) for row in sheet_rows(workbook, "Metadata")[1:] if row[0]}
    sectors = {}
    for column, code in enumerate(codes[3:], start=3):
        if not code:
            continue
        producer = float(country[column])
        ratio = float(conversion[code])
        index = float(price_index[code])
        sectors[code] = {
            "name": names[code],
            "kgPerUsdProducer2023": round(producer, 6),
            "purchaserRatio": round(ratio, 6),
            "priceIndex": index,
            "kgPerInr": round(producer * ratio * 100 / index / inr_per_usd, 8),
        }
    return sectors, inr_per_usd


def normalize(term):
    return " ".join(term.lower().replace(" and ", " & ").split())


def check_rules(rules, sectors):
    errors = []
    seen = {}
    for rule in rules:
        targets = [rule] + rule.get("byTitle", [])
        for target in targets:
            if target["sector"] is not None and target["sector"] not in sectors:
                errors.append(f"{rule['label']}: unknown sector {target['sector']}")
        for term in rule["crumbs"]:
            if normalize(term) != term:
                errors.append(f"{rule['label']}: write '{term}' as '{normalize(term)}'")
            # The same term may appear once generally and once under a department.
            key = (term, tuple(rule.get("under", [])))
            if key in seen:
                errors.append(f"'{term}' is in both '{seen[key]}' and '{rule['label']}'")
            seen[key] = rule["label"]
    return errors


def main():
    from openpyxl import load_workbook

    if not WORKBOOK.exists():
        print("Downloading Open CEDA 2025 from the AWS Open Data bucket...")
        download()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sectors, inr_per_usd = india_factors(workbook)

    rules = json.loads(RULES.read_text(encoding="utf-8"))["rules"]
    errors = check_rules(rules, sectors)
    if errors:
        print("\n".join(errors))
        sys.exit(1)

    used = {target["sector"] for rule in rules for target in [rule] + rule.get("byTitle", []) if target["sector"]}
    output = {
        "dataset": {**DATASET, "inrPerUsd": inr_per_usd, "built": date.today().isoformat()},
        "sectors": {code: sectors[code] for code in sorted(used)},
        "rules": rules,
    }
    OUTPUT.write_text(json.dumps(output, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(used)} sectors and {len(rules)} rules to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
